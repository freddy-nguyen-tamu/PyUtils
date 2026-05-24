from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont
from pathlib import Path

OUT_DIR = Path.cwd() / "pptx_30_slide_parts"
SLIDES_PER_PART = 30
DUPLICATES = 5
PLACEHOLDER_IMG = OUT_DIR / "placeholder.png"

OUT_DIR.mkdir(exist_ok=True)


def create_placeholder_image():
    if PLACEHOLDER_IMG.exists():
        return

    img = Image.new("RGB", (800, 450), "#D9EAF7")
    draw = ImageDraw.Draw(img)

    draw.rectangle([4, 4, 796, 446], outline="#4F81BD", width=8)
    draw.line(
        [(120, 330), (300, 190), (430, 310), (560, 160), (700, 330)],
        fill="#4F81BD",
        width=5,
    )
    draw.ellipse([575, 50, 665, 140], outline="#4F81BD", width=5)

    try:
        font = ImageFont.truetype("arial.ttf", 48)
    except Exception:
        font = ImageFont.load_default()

    text = "IMAGE PLACEHOLDER"
    bbox = draw.textbbox((0, 0), text, font=font)
    x = (800 - (bbox[2] - bbox[0])) / 2
    draw.text((x, 380), text, fill="#1F4E79", font=font)

    img.save(PLACEHOLDER_IMG)


def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_text(slide, text, x, y, w, h, size=18, bold=False):
    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    tf = box.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold

    return box


def add_title(slide):
    add_text(
        slide,
        "Presentation Title",
        0.6,
        0.35,
        12.1,
        0.6,
        size=30,
        bold=True,
    )


def add_small_text(slide, idx):
    add_text(
        slide,
        "Small supporting text",
        0.8,
        1.25 + idx * 0.42,
        5.8,
        0.25,
        size=11,
    )


def add_body(slide, idx, bullets=False):
    box = slide.shapes.add_textbox(
        Inches(0.8),
        Inches(1.25 + idx * 0.75),
        Inches(6.2),
        Inches(0.65),
    )
    tf = box.text_frame
    tf.clear()

    if bullets:
        bullet_texts = [
            "First bullet point",
            "Second bullet point",
        ]

        for i, txt in enumerate(bullet_texts):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = txt
            p.level = 0
            p.font.size = Pt(18)
    else:
        p = tf.paragraphs[0]
        p.text = "Main body text"
        p.font.size = Pt(18)


def add_image(slide, idx):
    slide.shapes.add_picture(
        str(PLACEHOLDER_IMG),
        Inches(7.5),
        Inches(1.25 + idx * 0.45),
        width=Inches(4.8),
        height=Inches(2.6),
    )


def add_case_slide(prs, spec):
    slide = blank_slide(prs)

    if spec["title_count"] == 1:
        add_title(slide)

    for i in range(spec["small_text_count"]):
        add_small_text(slide, i)

    for i in range(spec["body_count"]):
        add_body(slide, i, spec["bullets"])

    for i in range(spec["image_count"]):
        add_image(slide, i)


def is_all_zero_case(spec):
    return (
        spec["title_count"] == 0
        and spec["small_text_count"] == 0
        and spec["body_count"] == 0
        and spec["image_count"] == 0
    )


def generate_cases():
    # Pattern 1:
    # title + small text
    for title_count in [0, 1]:
        for small_text_count in range(10):
            spec = {
                "title_count": title_count,
                "small_text_count": small_text_count,
                "body_count": 0,
                "image_count": 0,
                "bullets": False,
            }

            if not is_all_zero_case(spec):
                yield spec

    # Pattern 2:
    # title + body text
    for title_count in [0, 1]:
        for body_count in range(10):
            spec = {
                "title_count": title_count,
                "small_text_count": 0,
                "body_count": body_count,
                "image_count": 0,
                "bullets": False,
            }

            if not is_all_zero_case(spec):
                yield spec

    # Pattern 3:
    # title + body text + image
    for title_count in [0, 1]:
        for body_count in range(10):
            for image_count in range(10):
                spec = {
                    "title_count": title_count,
                    "small_text_count": 0,
                    "body_count": body_count,
                    "image_count": image_count,
                    "bullets": False,
                }

                if not is_all_zero_case(spec):
                    yield spec

    # Pattern 4:
    # title + body text with 2 bullet points
    for title_count in [0, 1]:
        for body_count in range(10):
            spec = {
                "title_count": title_count,
                "small_text_count": 0,
                "body_count": body_count,
                "image_count": 0,
                "bullets": True,
            }

            if not is_all_zero_case(spec):
                yield spec

    # Pattern 5:
    # title + body text with 2 bullet points + image
    for title_count in [0, 1]:
        for body_count in range(10):
            for image_count in range(10):
                spec = {
                    "title_count": title_count,
                    "small_text_count": 0,
                    "body_count": body_count,
                    "image_count": image_count,
                    "bullets": True,
                }

                if not is_all_zero_case(spec):
                    yield spec


def save_part(prs, part, slide_count):
    out = OUT_DIR / f"part_{part:03}.pptx"
    prs.save(out)
    print(f"Wrote {out} with {slide_count} slides")


def main():
    create_placeholder_image()

    prs = create_presentation()
    part = 1
    slide_count = 0
    total_slides = 0
    total_cases = 0

    for spec in generate_cases():
        total_cases += 1

        for _ in range(DUPLICATES):
            if slide_count == SLIDES_PER_PART:
                save_part(prs, part, slide_count)
                part += 1
                slide_count = 0
                prs = create_presentation()

            add_case_slide(prs, spec)

            slide_count += 1
            total_slides += 1

    if slide_count > 0:
        save_part(prs, part, slide_count)

    print("Done.")
    print(f"Total unique cases: {total_cases}")
    print(f"Total slides: {total_slides}")
    print(f"Output folder: {OUT_DIR}")
    print(f"Placeholder image: {PLACEHOLDER_IMG}")


if __name__ == "__main__":
    main()